"""Word v3 y diagramas derivados del modelo actual. No requiere conexión a la BD.

Conserva el original y v2; no contiene ni consulta credenciales para el documento.
"""
import json
import math
import re
import zipfile
from datetime import datetime
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from sqlalchemy.dialects import postgresql
from database import metadata

BASE = Path(__file__).resolve().parent
OUT = BASE / 'docs_generados'
ORIGINAL = BASE.parent / 'Documentacion Tecnica Chatbot Bitel.docx'
DESTINO = BASE.parent / 'Documentacion Tecnica Chatbot Bitel v3.docx'
NAVY = '1F4E79'
MD = []

TABLAS = {
    'clientes': 'Titulares comerciales identificados por DNI único, independientes de Telegram. '
                'El alta por el bot crea el titular y su primera línea en una sola transacción.',
    'contactos': 'Participantes de Telegram. Su id es el chat_id. Escribir al bot crea un contacto, '
                 'no un cliente comercial. No guarda un DNI permanente.',
    'asesores': 'IDs Telegram autorizados para atención humana; la lista operativa proviene de ADMIN_CHAT_IDS.',
    'conversaciones': 'Historial de conversaciones por chat, abiertas, cerradas o reiniciadas. '
                      'No depende de que el visitante haya proporcionado un DNI.',
    'mensajes': 'Textos de entrada/salida, avisos de asesores y marcas de tiempo por conversación.',
    'sesiones': 'Espejo del estado temporal de un chat en JSON. No es una cuenta de cliente; '
                'se vacía al reiniciar y no se restaura la identificación anterior.',
    'asignaciones': 'Historial de asesores que tomaron un chat. fin nulo indica una asignación sin cerrar. '
                    'La exclusividad de una toma se controla en la instancia del bot.',
    'planes': 'Catálogo único de planes: se utiliza tanto para consulta informativa como para adquirir líneas.',
    'equipos': 'Catálogo de equipos y precios; su adquisición registra una solicitud para atención.',
    'catalogo': 'Promociones, servicios y paquetes. requiere_direccion decide si un servicio pide '
                'dirección de instalación; en caso contrario pide una línea asociada.',
    'lineas': 'Números contratados por un cliente. Un DNI puede tener varios números; cada línea tiene '
              'un plan. numero es único. activa se conserva en el esquema pero no filtra las autogestiones.',
    'deudas': 'Importes por línea, no por chat ni globalmente por cliente. Solo pagada=false se muestra '
              'como pendiente. El bot no genera deudas al contratar.',
    'solicitudes': 'Gestiones comerciales con cliente titular, datos, estado, chat de origen y UUID '
                   'de operación para impedir repetir la misma confirmación.',
    'historial_solicitudes': 'Cambios de estado de una solicitud y responsable. Se registra junto a '
                            'la solicitud en la misma transacción.'
}
CAMPO = {
    'dni': 'Documento de ocho dígitos; conservar ceros iniciales. UNIQUE y CHECK de formato.',
    'telefono': 'Teléfono de contacto de nueve dígitos comenzando por 9; no es la línea asignada.',
    'chat_id': 'Canal Telegram de interacción o de origen; no vincula permanentemente el DNI a la cuenta.',
    'cliente_id': 'ID interno del titular comercial, independiente del ID Telegram.',
    'numero': 'Número de nueve dígitos comenzando por 9, asignado sin repetir otro número de lineas.',
    'precio': 'Importe decimal en soles. Se vuelve a comprobar al confirmar una adquisición.',
    'importe': 'Monto de la deuda en soles; precisión NUMERIC(12,2).',
    'vencimiento': 'Fecha de vencimiento guardada como texto YYYY-MM-DD; no es una columna DATE.',
    'pagada': 'Indica si la deuda está pagada; los registros pagados no suman al total pendiente.',
    'activa': 'Campo histórico; se conserva sin considerar su estado en estos flujos.',
    'datos': 'JSON con estado del chat o detalle del formulario, según la tabla.',
    'operacion': 'UUID de una confirmación. UNIQUE garantiza que el reintento no repita el registro.',
    'asesor_id': 'Asesor responsable de una asignación o solicitud.',
    'responsable_id': 'ID del participante que realizó el cambio; no es una FK por abarcar cliente/asesor.',
    'autor_id': 'ID del autor del mensaje; puede ser nulo en respuestas del bot.',
    'rol': 'Origen del texto: cliente (participante), bot o asesor.',
    'categoria': 'PROMOCIONES, SERVICIOS o PAQUETES.',
    'requiere_direccion': 'Solicita dirección de instalación para este servicio.',
    'conversacion_id': 'Conversación a la que pertenece el mensaje o la asignación.',
    'solicitud_id': 'Solicitud cuyo cambio de estado se registra.',
    'plan_id': 'Plan contratado para esta línea.',
    'linea_id': 'Línea a la que pertenece la deuda.',
    'tipo': 'linea_nueva, portabilidad, baja, equipo o servicio. Consultar deuda no crea solicitud.',
    'nombre': 'Nombre completo del cliente, nombre de producto o nombre mostrado por Telegram.',
    'descripcion': 'Contenido informativo consultado directamente en PostgreSQL.',
    'creado': 'Fecha de alta con zona horaria.', 'fecha': 'Fecha de registro con zona horaria.',
    'inicio': 'Inicio de la conversación/asignación.', 'fin': 'Cierre; nulo mientras no haya finalizado.',
}


def font(size, bold=False):
    candidates = [Path('C:/Windows/Fonts') / ('arialbd.ttf' if bold else 'arial.ttf'),
                  Path('/usr/share/fonts/truetype/dejavu') / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default(size=size)


def arrow(draw, a, b, fill='#64748B'):
    draw.line([a, b], fill=fill, width=4)
    angle = math.atan2(b[1]-a[1], b[0]-a[0])
    points = [b] + [(b[0]-20*math.cos(angle+d), b[1]-20*math.sin(angle+d)) for d in (-.4, .4)]
    draw.polygon(points, fill=fill)


def boundary(rect, other):
    x = (rect[0]+rect[2])/2
    y = (rect[1]+rect[3])/2
    dx, dy = other[0]-x, other[1]-y
    times = []
    if dx:
        times.append(((rect[2] if dx > 0 else rect[0])-x)/dx)
    if dy:
        times.append(((rect[3] if dy > 0 else rect[1])-y)/dy)
    t = min(times)
    return x+dx*t, y+dy*t


def er_diagram(filename, positions, title):
    """Solo PK/FK en la figura; el diccionario Word detalla todos los campos."""
    rects = {}
    fields = {}
    for name, (x, y) in positions.items():
        fields[name] = [c for c in metadata.tables[name].columns if c.primary_key or c.foreign_keys]
        if name == 'clientes':
            fields[name] += [metadata.tables[name].c.dni]
        if name == 'lineas':
            fields[name] += [metadata.tables[name].c.numero]
        rects[name] = (x, y, x+620, y+85+len(fields[name])*52)
    w = max(r[2] for r in rects.values())+90
    h = max(r[3] for r in rects.values())+125
    image = Image.new('RGB', (w, h), 'white')
    d = ImageDraw.Draw(image)
    d.text((50, 20), title, font=font(36, True), fill='#1F4E79')
    for name in positions:
        for fk in metadata.tables[name].foreign_keys:
            parent = fk.column.table.name
            if parent not in positions:
                continue
            childrect, parentrect = rects[name], rects[parent]
            ca = ((childrect[0]+childrect[2])/2, (childrect[1]+childrect[3])/2)
            cb = ((parentrect[0]+parentrect[2])/2, (parentrect[1]+parentrect[3])/2)
            arrow(d, boundary(childrect, cb), boundary(parentrect, ca))
    for name, r in rects.items():
        x, y, right, bottom = r
        d.rounded_rectangle(r, radius=12, fill='#EFF6FF', outline='#1F4E79', width=3)
        d.rectangle((x, y, right, y+68), fill='#1F4E79')
        f = font(34, True)
        if d.textlength(name, font=f) > 585:
            f = font(29, True)
        d.text((x+16, y+13), name, font=f, fill='white')
        for i, c in enumerate(fields[name]):
            tags = (['PK'] if c.primary_key else []) + (['FK'] if c.foreign_keys else [])
            if c.unique:
                tags.append('UQ')
            d.text((x+15, y+80+i*52), c.name + ' [' + ', '.join(tags) + ']', font=font(30), fill='#172554')
    d.text((50, h-75), 'PK: clave primaria · FK: clave foránea · UQ: único · Flecha: FK → tabla referenciada',
           font=font(24), fill='#475569')
    image.save(OUT / filename)


def flow_diagram(filename, title, labels):
    image = Image.new('RGB', (1900, 160+len(labels)*150), 'white')
    d = ImageDraw.Draw(image)
    d.text((45, 25), title, font=font(38, True), fill='#1F4E79')
    for i, label in enumerate(labels):
        y = 110+i*150
        d.rounded_rectangle((100, y, 1800, y+105), radius=16, fill='#EFF6FF', outline='#1F4E79', width=3)
        d.text((135, y+30), label, font=font(30), fill='#172554')
        if i < len(labels)-1:
            arrow(d, (950, y+105), (950, y+150))
    image.save(OUT / filename)


def mermaid():
    lines = ['erDiagram']
    for table in metadata.sorted_tables:
        lines += [f'    {table.name} {{']
        for c in table.columns:
            type_ = re.sub(r'[^a-zA-Z0-9_]', '_', str(c.type)).strip('_')
            tags = ','.join((['PK'] if c.primary_key else []) + (['FK'] if c.foreign_keys else []) + (['UK'] if c.unique else []))
            lines.append(f'        {type_} {c.name}' + (f' {tags}' if tags else ''))
        lines += ['    }']
    for table in metadata.sorted_tables:
        for fk in sorted(table.foreign_keys, key=lambda fk: fk.parent.name):
            parent = 'o|' if fk.parent.nullable else '||'
            child = 'o|' if fk.parent.primary_key or fk.parent.unique else 'o{'
            lines.append(f'    {fk.column.table.name} {parent}--{child} {table.name} : {fk.parent.name}')
    (OUT / 'modelo_datos.mmd').write_text('\n'.join(lines)+'\n', encoding='utf-8')


def p(doc, text, style=None):
    paragraph = doc.add_paragraph(text, style)
    MD.append(text + '\n')
    return paragraph


def heading(doc, text, level=1):
    doc.add_heading(text, level)
    MD.append('#'*level + ' ' + text + '\n')


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    MD.append('| ' + ' | '.join(headers) + ' |\n| ' + ' | '.join(['---']*len(headers)) + ' |')
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), NAVY)
        cell._tc.get_or_add_tcPr().append(shd)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        MD.append('| ' + ' | '.join(str(v).replace('\n', '<br>').replace('|', '/') for v in values) + ' |')
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)
    MD.append('')
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    return t


def picture(doc, filename, caption):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_picture(str(OUT / filename), width=Cm(16.5))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = p(doc, caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    MD.append(f'![{caption}]({filename})\n')


def build():
    OUT.mkdir(exist_ok=True)
    MD.clear()
    report = json.loads((BASE / 'resultados_verificacion.json').read_text(encoding='utf-8'))
    doc = Document(str(ORIGINAL)) if ORIGINAL.exists() else Document()
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
        section.footer.paragraphs[0].text = 'TeleBOT Integrador · Documentación técnica v3.0'
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3'):
        doc.styles[name].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string(NAVY)

    flow_diagram('figura1_arquitectura_v3.png', 'Arquitectura y persistencia', [
        'Cliente o asesor → Telegram Bot API (long polling HTTPS)',
        'Bot Python → control de acceso → máquina de estados',
        'Consultas informativas → catálogos PostgreSQL',
        'Autogestiones → DNI → cliente → líneas → deuda / contratación',
        'Repositorio transaccional → PostgreSQL 17 (volumen persistente)',
        'pgAdmin → inspección administrativa de tablas y relaciones'])
    flow_diagram('figura2_flujo_v3.png', 'Entrada intuitiva y autogestión', [
        'Cualquier primer texto → bienvenida + menú (sin escribir /start)',
        '1–4: consulta libre · 5: asesor · 6: autogestiones',
        'Elegir gestión → COM_DNI → buscar cliente por DNI',
        'DNI existente: continuar · DNI nuevo: ofrecer primera línea',
        'Elegir plan/producto o teléfono; recoger datos necesarios',
        'Resumen → confirmar / cancelar → registro en PostgreSQL',
        'Deuda: seleccionar teléfono → pendientes y total de ese número'])
    er_diagram('figura3_comercial_v3.png', {
        'clientes': (50, 110), 'planes': (820, 110), 'catalogo': (1590, 110),
        'solicitudes': (50, 560), 'lineas': (820, 560), 'equipos': (1590, 560),
        'historial_solicitudes': (50, 1060), 'deudas': (820, 1060)
    }, 'Modelo comercial: DNI → varias líneas → deudas por número')
    er_diagram('figura4_atencion_v3.png', {
        'contactos': (50, 110), 'asesores': (1300, 110),
        'conversaciones': (50, 480), 'asignaciones': (1300, 480),
        'mensajes': (50, 950), 'sesiones': (1300, 950)
    }, 'Historial de atención por chat, separado del cliente por DNI')
    mermaid()

    heading(doc, 'DOCUMENTACIÓN TÉCNICA')
    p(doc, 'Chatbot de atención y autogestiones telefónicas — TeleBOT Integrador')
    table(doc, ['Dato', 'Detalle'], [
        ('Autores', 'Vidal Maza. J, Zevallos Miranda. C, Sarmiento Castro. L.'),
        ('Versión', '3.0 — autogestiones por DNI'), ('Fecha de generación', datetime.now().strftime('%d/%m/%Y')),
        ('Estado', 'Implementado sobre PostgreSQL; Telegram simulado en pruebas automatizadas.'),
        ('Alcance', 'Consulta pública, identificación por DNI, múltiples líneas, deuda por teléfono y solicitudes de servicio.')])
    p(doc, 'Se conservan el documento original y la versión 2.0. Esta versión reemplaza la descripción '
           'de clientes por Telegram y las altas demo automáticas por el modelo comercial por DNI.')
    titles = ['Descripción general del proyecto', 'Alcance actual', 'Arquitectura y componentes',
              'Tecnologías y dependencias', 'Configuración y estructura', 'Conversación y estados',
              'Catálogos desde PostgreSQL', 'Gestión de asesores', 'Requerimientos funcionales',
              'Requerimientos no funcionales', 'Instalación y operación', 'Pruebas y resultados',
              'Manejo de errores e identificación', 'Limitaciones actuales', 'Base de datos y diccionario',
              'Diagrama de la base de datos']
    heading(doc, 'Índice')
    for i, title in enumerate(titles, 1):
        p(doc, f'{i}. {title}')
    doc.add_page_break()

    heading(doc, '1. Descripción general del proyecto')
    p(doc, 'Bot de Telegram para consultar información telefónica y ejecutar autogestiones sobre '
           'registros PostgreSQL. Cualquier primer mensaje de texto muestra bienvenida y menú. '
           'Las consultas de catálogo no requieren DNI; cada nueva gestión comercial solicita un DNI '
           'y busca al titular. No existe una vinculación permanente entre DNI y cuenta Telegram.')
    p(doc, 'Un cliente tiene un DNI único y puede tener varias líneas, cada una con número único y plan. '
           'Confirmar una nueva línea guarda el registro en la base; un alta de cliente se confirma junto '
           'a su primera línea. La deuda se consulta por teléfono, nunca como un importe inventado.')

    heading(doc, '2. Alcance actual')
    table(doc, ['Función', 'Resultado implementado'], [
        ('Consulta', 'Planes, promociones, servicios y paquetes desde PostgreSQL, sin identificación.'),
        ('Línea nueva', 'Cliente existente: línea adicional. Cliente nuevo: titular y primera línea en una transacción.'),
        ('Deuda', 'DNI → seleccionar teléfono propio → registros no pagados y total.'),
        ('Portabilidad', 'Plan, número de origen y operador; solicitud pendiente para asesor.'),
        ('Baja', 'Número asociado, motivo y detalle si es Otro; solicitud pendiente.'),
        ('Equipo', 'Equipo, contado/evaluación de cuotas y dirección de entrega; solicitud pendiente.'),
        ('Servicio', 'Servicio, dirección de instalación o línea asociada; solicitud pendiente.'),
        ('Mis solicitudes', 'Últimas 20 del cliente indicado por DNI, desde cualquier chat.'),
        ('Asesor', 'Toma exclusiva, respuesta, cierre y seguimiento de solicitudes.')])

    heading(doc, '3. Arquitectura y componentes')
    picture(doc, 'figura1_arquitectura_v3.png', 'Figura 1. Arquitectura del sistema actual.')
    table(doc, ['Componente', 'Responsabilidad'], [
        ('bot_corregido.py', 'Acceso público/privado, menú, handlers, asesor y catálogo informativo.'),
        ('commerce.py', 'DNI, captura de datos, teléfono seleccionado, resumen y confirmación.'),
        ('database.py', 'Modelo y repositorio: transacciones, bloqueo por DNI, unicidad e idempotencia.'),
        ('persistence.py', 'Registra entradas, respuestas y sesiones fuera del bucle asíncrono.'),
        ('PostgreSQL / pgAdmin', 'Datos persistentes y administración web.'),
        ('Alembic', 'Migraciones 001 histórica y 002 de identidad por DNI.')])

    heading(doc, '4. Tecnologías y dependencias')
    dependencies = [line.strip() for line in (BASE / 'requirements.txt').read_text().splitlines() if line.strip()]
    table(doc, ['Tecnología', 'Versión / definición'], [('Python', '3.14'), ('PostgreSQL', '17'),
        ('pgAdmin', '9'), ('Docker Compose', 'Servicios bot, postgres y pgadmin')] +
        [(line.split('==')[0], line.split('==')[-1]) for line in dependencies])
    p(doc, 'python-docx y Pillow se instalan separadamente desde requirements-docs.txt para regenerar el Word y los PNG.')

    heading(doc, '5. Configuración y estructura')
    table(doc, ['Variable .env', 'Uso'], [
        ('TELEGRAM_BOT_TOKEN', 'Token de BotFather; no se publica en documentos.'),
        ('ACCESO_PUBLICO', 'True: cualquier chat privado. False: solo IDs autorizados y asesores.'),
        ('ADMIN_CHAT_IDS', 'IDs positivos de asesores, separados por comas.'),
        ('ALLOWED_USER_IDS', 'Lista de acceso para modo privado.'),
        ('POSTGRES_PASSWORD / DB_APP_PASSWORD / PGADMIN_PASSWORD', 'Contraseñas independientes de administrador, aplicación y web.')])
    p(doc, 'Compose solo entrega las variables necesarias a cada contenedor. .env y respaldos quedan '
           'fuera de la imagen y de Git. Tras cambiar IDs o acceso: docker compose up -d --force-recreate bot.')
    p(doc, 'Archivos adicionales: cargar_ejemplos.py, test_flujo.py, test_postgres.py, benchmark.py, '
           'resultados_verificacion.json, generar_documentacion.py, docs_generados/ y respaldos/.')

    heading(doc, '6. Conversación y estados')
    picture(doc, 'figura2_flujo_v3.png', 'Figura 2. Entrada, identificación y confirmación.')
    table(doc, ['Estado', 'Función'], [
        ('MENU', '1–4 catálogos, 5 asesor, 6 autogestiones.'),
        ('PLANES / PROMOCIONES / SERVICIOS / PAQUETES', 'Opciones obtenidas de PostgreSQL; 0 vuelve al menú.'),
        ('OTRA_CONSULTA', '1 nueva consulta, 2 finalizar, 3 regresar.'),
        ('ASESOR', 'Enrutamiento a asesor asignado o aviso a los disponibles.'),
        ('GESTIONES', '1 línea, 2 portabilidad, 3 baja, 4 equipo, 5 deuda, 6 solicitudes, 7 servicio.'),
        ('COM_DNI / COM_ALTA', 'Validación de DNI; si no existe, ofrecer primera línea sin crear cliente aún.'),
        ('COM_PRODUCTO / COM_LINEA', 'Elegir producto o un número asociado al DNI.'),
        ('COM_NOMBRE / COM_CONTACTO', 'Nombre completo y teléfono de contacto del cliente nuevo.'),
        ('COM_NUMERO / COM_OPERADOR', 'Número a portar y operador de origen.'),
        ('COM_MOTIVO / COM_DETALLE', 'Motivo de baja y descripción para Otro.'),
        ('COM_MODALIDAD / COM_DIRECCION', 'Forma de adquisición y ubicación de entrega/instalación.'),
        ('COM_CONFIRMAR', 'Resumen de datos: 1 confirmar, 2 cancelar, 0 menú.')])
    p(doc, 'menú, menu o inicio reinicia desde texto; /start sigue como comando opcional. Reiniciar '
           'cancela el formulario y la atención de chat, pero no solicitudes confirmadas. El DNI se '
           'solicita de nuevo al comenzar cada gestión. No se requiere DNI para hablar con un asesor.')

    heading(doc, '7. Catálogos desde PostgreSQL')
    p(doc, 'No existe un catálogo comercial de respaldo en Python. PLANES consulta planes; los demás '
           'menús filtran catalogo por categoría. Equipos se obtiene de equipos. Un catálogo vacío se '
           'informa como tal. La carga inicial de ejemplos es explícita; no se restaura al reiniciar.')
    p(doc, 'En una adquisición se conserva el precio mostrado y se comprueba nuevamente al confirmar. '
           'Si cambió o el producto no existe, se pide volver al menú y seleccionar de nuevo.')

    heading(doc, '8. Gestión de asesores')
    table(doc, ['Comando', 'Uso'], [('/id', 'ID Telegram del participante autorizado por el modo de acceso.'),
        ('/tomar ID_CHAT', 'Tomar un chat pendiente; primer asesor disponible.'),
        ('/reply ID_CHAT mensaje', 'Responder únicamente al chat propio.'),
        ('/finalizar ID_CHAT', 'Cerrar únicamente el chat propio.'),
        ('/solicitudes', 'Últimas 20 solicitudes globales.'),
        ('/estado ID en_atencion|completada|cancelada', 'Transiciones válidas y exclusividad del asesor responsable.')])
    p(doc, 'pendiente → en_atencion → completada/cancelada, o pendiente → cancelada. La línea nueva '
           'se registra directamente como completada. Las otras gestiones se registran pendientes. '
           'El chat de origen recibe avisos; cliente_id nunca se usa como destino Telegram.')

    heading(doc, '9. Requerimientos funcionales')
    requirements = [
        ('RF01', 'Iniciar con cualquier primer texto; /start es opcional.'),
        ('RF02', 'Bienvenida y menú principal con seis opciones.'),
        ('RF03', 'Opciones numéricas con cancelación intuitiva por texto.'),
        ('RF04–RF07', 'Consultar planes, promociones, servicios y paquetes desde PostgreSQL.'),
        ('RF08', 'Entradas inválidas no avanzan el formulario.'),
        ('RF09–RF11', 'Nueva consulta, regreso y cierre de la consulta automática.'),
        ('RF12–RF16', 'Derivación, avisos, toma exclusiva, respuesta y cierre por asesor.'),
        ('RF17', 'Obtener ID Telegram según control de acceso.'),
        ('RF18', 'Adquirir línea con DNI, plan y número único.'),
        ('RF19', 'Solicitar portabilidad con número y operador.'),
        ('RF20', 'Solicitar baja de un número asociado con motivo.'),
        ('RF21', 'Solicitar equipo con modalidad y dirección.'),
        ('RF22', 'Elegir teléfono y consultar exclusivamente su deuda pendiente.'),
        ('RF23–RF24', 'Consultar solicitudes y cambiar estados con historial.'),
        ('RF25', 'Persistir interacciones sin crear clientes comerciales por escribir al bot.'),
        ('RF26', 'Acceso público o privado configurable; asesor siempre restringido.'),
        ('RF27', 'Identificar por DNI sin vincularlo permanentemente a Telegram.'),
        ('RF28', 'Ofrecer alta ante DNI desconocido; cliente y primera línea atómicos.'),
        ('RF29', 'Un DNI admite varios números únicos con sus respectivos planes.'),
        ('RF30', 'Servicio pide dirección o línea según datos del catálogo.'),
        ('RF31', 'Cancelación no registra cliente/solicitud; reintento de la misma confirmación no duplica.')]
    table(doc, ['Código', 'Requerimiento actualizado'], requirements)

    heading(doc, '10. Requerimientos no funcionales')
    table(doc, ['Código', 'Atributo y evidencia'], [
        ('RNF01', 'Usabilidad: inicio por texto, preguntas pertinentes y resumen legible.'),
        ('RNF02', 'Compatibilidad: Python y contenedores Linux.'),
        ('RNF03', 'Rendimiento: p95 local < 1000 ms; ver medición y alcance en sección 12.'),
        ('RNF04', 'Confiabilidad: separación por chat, transacciones de alta y deuda por línea.'),
        ('RNF05–RNF06', 'Secretos fuera del código; permisos de asesor independientes del acceso público.'),
        ('RNF07–RNF09', 'Separación de módulos, validaciones y dependencias fijadas.'),
        ('RNF10', 'Múltiples participantes y asesores en una instancia; sin replicación.'),
        ('RNF11–RNF12', 'Persistencia PostgreSQL e historial de estados.'),
        ('RNF13', 'Unicidad de DNI/número y confirmación idempotente bajo concurrencia.')])

    heading(doc, '11. Instalación y operación')
    for command in ['docker compose up -d --build', 'docker compose ps', 'docker compose logs --tail 50 bot',
                    'docker compose stop bot', 'docker compose start bot']:
        p(doc, command)
    p(doc, 'pgAdmin: http://localhost:5050, admin@telebot.com. PostgreSQL: telebot; aplicación: '
           'telebot_app. Contraseñas desde .env. No usar down -v para detener: elimina volúmenes.')
    heading(doc, '11.1 Migración y carga autorizada', 2)
    p(doc, 'La revisión 002 retira las tablas comerciales demo anteriores y conserva el historial de '
           'atención, renombrando clientes anterior a contactos. Es una migración específica del avance '
           'autorizado, no una limpieza en cada arranque. Respaldo: respaldos/telebot_antes_dni.dump.')
    p(doc, 'En instalación nueva: python -m alembic upgrade head y luego cargar_ejemplos.py dentro '
           'del contenedor. El script carga 15 clientes y 30 líneas en una transacción y no duplica '
           'la carga existente. No repara filas de ejemplo eliminadas ni sustituye otros catálogos.')
    p(doc, 'DNI de prueba: 00001001–00001015. Cinco clientes con una línea, cinco con dos y cinco '
           'con tres. Hay números sin deuda, con deuda pagada y con dos pendientes. 00001011 permite '
           'probar tres números y dos deudas pendientes en el tercero. Las adquisiciones no generan deuda.')

    heading(doc, '12. Pruebas y resultados')
    p(doc, 'Las pruebas usan PostgreSQL real y Telegram simulado en telebot_test. La medición '
           'recrea los ejemplos solo en esa base. Ninguna prueba envía mensajes a cuentas reales.')
    table(doc, ['Resultado', 'Valor'], [(k, v) for k, v in report['pruebas'].items()])
    if report.get('rendimiento'):
        table(doc, ['Métrica local', 'Resultado'], [(k, v) for k, v in report['rendimiento'].items()])
    p(doc, 'Cobertura: consulta pública sin alta, 15 clientes/30 líneas, deuda por número, '
           'DNI inválido, cancelación, alta atómica, líneas adicionales desde otro chat, colisión de '
           'números, confirmación concurrente, rollback, cambio de precio, portabilidad/equipo/servicio/'
           'baja y seguimiento por asesor. El benchmark secuencial no mide la red Telegram.')
    p(doc, 'Comando: docker compose run --rm --no-deps -e DB_NAME=telebot_test bot python -m unittest -v test_flujo test_postgres')
    p(doc, 'Medición: docker compose run --rm --no-deps -e DB_NAME=telebot_test bot python benchmark.py')

    heading(doc, '13. Manejo de errores e identificación')
    p(doc, 'DNI: ocho dígitos ASCII; teléfono: nueve comenzando por 9. Nombre completo, operador y '
           'dirección tienen validaciones de longitud. Si falla PostgreSQL se avisa sin continuar el '
           'flujo. Si falla la asignación de número, se revierte también el alta del cliente.')
    p(doc, 'Cada confirmación tiene un UUID único. Un bloqueo transaccional por DNI serializa '
           'confirmaciones simultáneas; UNIQUE(numero) más reintentos limita las colisiones. Se verifica '
           'el producto/precio y la pertenencia de la línea en la base, no solo en el texto recibido.')
    p(doc, 'Por decisión del diseño, un DNI registrado se puede usar desde cualquier chat. Su existencia '
           'es identificación del registro, no autenticación del titular. No hay OTP ni validación externa '
           'del DNI. El teléfono solicitado es un dato de contacto, no una verificación de identidad.')

    heading(doc, '14. Limitaciones actuales')
    for text in [
        'Los 15 clientes y el catálogo inicial son ejemplos. El bot trabaja con registros de la base; no genera importes al consultar.',
        'Asignar un número en esta base no activa una línea en una operadora. No hay integración Bitel ni cobros.',
        'Portabilidad, bajas, equipos y servicios son solicitudes. Cambiar su estado no ejecuta la operación externa.',
        'No hay turnos ni panel operativo de asesores; pgAdmin es administración técnica.',
        'Sesiones se reinician al arrancar. Updates pendientes de Telegram se descartan deliberadamente.',
        'No hay tabla audit_log adicional ni transacción distribuida entre Telegram y PostgreSQL.',
        'producto_id y linea_id dentro del JSON no son FKs; contract verifica producto y pertenencia de línea.',
        'Capturas de Telegram y atención con dos cuentas reales están pendientes de evidencia manual.'
    ]:
        p(doc, text, 'List Bullet')

    heading(doc, '15. Base de datos y diccionario')
    p(doc, '14 tablas de aplicación más alembic_version. Este diccionario se deriva directamente '
           'del modelo SQLAlchemy, incluyendo tipos, nulabilidad, PK, FK y unicidad; no inventa '
           'columnas ni relaciones. Los valores default de SQLAlchemy pueden aplicarse desde Python '
           'y no ser defaults del servidor; se distinguen abajo.')
    for i, (name, description) in enumerate(TABLAS.items(), 1):
        heading(doc, f'15.{i} {name}', 2)
        p(doc, description)
        rows = []
        for c in metadata.tables[name].columns:
            rules = ['NULL' if c.nullable else 'NOT NULL']
            if c.primary_key:
                rules.append('PK')
            if c.unique:
                rules.append('UNIQUE')
            for fk in c.foreign_keys:
                rules.append('FK → '+fk.target_fullname)
            if c.server_default is not None:
                rules.append('Default BD: '+str(c.server_default.arg))
            elif c.default is not None:
                rules.append('Default Python: '+str(c.default.arg))
            detail = CAMPO.get(c.name, 'Identificador interno del registro.')
            if c.name == 'id' and name in ('contactos', 'asesores'):
                detail = 'ID Telegram, no DNI ni identificador comercial.'
            if c.name == 'estado':
                detail = ('abierta, cerrada o reiniciada.' if name == 'conversaciones'
                          else 'pendiente, en_atencion, completada o cancelada.')
            rows.append((c.name, str(c.type.compile(dialect=postgresql.dialect())), '; '.join(rules), detail))
        table(doc, ['Campo', 'Tipo PostgreSQL', 'Restricciones / default', 'Descripción'], rows)
    heading(doc, '15.15 alembic_version', 2)
    p(doc, 'Tabla técnica administrada por Alembic: version_num VARCHAR(32), PK NOT NULL. '
           'Guarda la revisión aplicada (002). No representa datos de clientes.')

    heading(doc, '16. Diagrama de la base de datos')
    picture(doc, 'figura3_comercial_v3.png', 'Figura 3. Dominio comercial. solicitudes.asesor_id referencia asesores en la figura 4.')
    picture(doc, 'figura4_atencion_v3.png', 'Figura 4. Atención por chat separada de la identidad por DNI.')
    table(doc, ['FK', 'Tabla referenciada'], [
        (f'{t.name}.{fk.parent.name}', fk.target_fullname)
        for t in metadata.sorted_tables for fk in sorted(t.foreign_keys, key=lambda f: f.parent.name)])
    p(doc, 'Relación principal: clientes 1:N lineas; planes 1:N lineas; lineas 1:N deudas. '
           'Cada contacto tiene 0..1 sesión y N conversaciones. solicitudes tiene chat_id como dato '
           'de origen sin FK; no existe FK permanente entre contactos y clientes. El Mermaid completo '
           'incluye todos los campos y relaciones en docs_generados/modelo_datos.mmd.')

    heading(doc, 'Anexo A. Flujo histórico original')
    if ORIGINAL.exists():
        with zipfile.ZipFile(ORIGINAL) as archive:
            if 'word/media/image3.png' in archive.namelist():
                (OUT / 'anexo_a_flujo_original.png').write_bytes(archive.read('word/media/image3.png'))
                picture(doc, 'anexo_a_flujo_original.png', 'Figura A1. Referencia histórica v1; no describe el alta por DNI actual.')
    heading(doc, 'Anexo B. Trazabilidad')
    table(doc, ['Grupo', 'Verificación'], [
        ('Entrada, catálogo, permisos y asesor', 'test_flujo.py + integración PostgreSQL.'),
        ('DNI, varias líneas, deuda y adquisiciones', 'test_postgres.py: consultas, transacciones y concurrencia.'),
        ('Persistencia y rendimiento', 'benchmark.py y resultados_verificacion.json.'),
        ('Esquema y relaciones', 'Metadatos SQLAlchemy y migraciones 001 → 002.')])
    doc.core_properties.title = 'TeleBOT Integrador — Autogestiones por DNI'
    doc.core_properties.subject = 'Documentación técnica v3.0'
    doc.save(DESTINO)
    (OUT / 'documentacion_v3.md').write_text('\n'.join(MD), encoding='utf-8')
    # Comprobar el artefacto guardado, no solo el objeto en memoria.
    with zipfile.ZipFile(DESTINO) as archive:
        assert archive.testzip() is None, 'DOCX corrupto'
    reopened = Document(DESTINO)
    headings = {paragraph.text for paragraph in reopened.paragraphs}
    assert all(f'15.{i} {name}' in headings for i, name in enumerate(TABLAS, 1))
    assert len(reopened.inline_shapes) >= 4
    print('Word:', DESTINO)
    print('Modelo: 14 tablas de aplicación + alembic_version; PNG y Mermaid en', OUT)
    print('Verificación DOCX: OK; diccionario completo y figuras presentes.')


if __name__ == '__main__':
    build()
